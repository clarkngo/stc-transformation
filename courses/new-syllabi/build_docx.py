#!/usr/bin/env python3
"""
Build the Word syllabi for AI 251, AI 320, CS 352, and AI 371 from their HTML pages.

    python3 courses/new-syllabi/build_docx.py            # all four
    python3 courses/new-syllabi/build_docx.py AI-320     # just one

The HTML page is the source of truth. Only the <article class="syllabus-doc"> element is
exported; the review bar, navigation, and anything else on the page stay web-only.
Styles, margins, and the page-number footer come from reference.docx, which is the
AI 410 syllabus (2026-09-14) with its body removed. The CityU banner is added at the top.

Needs pandoc (3.x) and python-docx.
"""

import re
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
COURSES = HERE.parent
COURSES_TO_BUILD = {
    "AI-251": "AI251-Syllabus-DRAFT.docx",
    "AI-320": "AI320-Syllabus-DRAFT.docx",
    "CS-352": "CS352-Syllabus-DRAFT.docx",
    "AI-371": "AI371-Syllabus-DRAFT.docx",
}

ARTICLE = re.compile(r'<article class="syllabus-doc"[^>]*>(.*?)</article>', re.S)


def to_word_html(article):
    """Map the page's heading levels onto the AI 410 docx: h1 title and h2 sections -> Heading 1,
    h3 -> Heading 3 (policy subsections), h4 -> Heading 4, h5 (the PeopleSoft intro) -> Heading 2,
    the shaded italic style AI 410 uses there."""
    html = article
    html = re.sub(r"<(/?)h2\b", r"<\1h1", html)
    html = re.sub(r"<(/?)h5\b", r"<\1h2", html)
    banner = (HERE / "cityu-syllabus-banner.png").as_posix()
    return f'<html><body><p><img src="{banner}" style="width:6.5in" alt="City University of Seattle"></p>{html}</body></html>'


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def polish(path):
    d = docx.Document(path)
    grid = d.styles["Table Grid"]
    for t in d.tables:
        t.style = grid
        t.autofit = True
        if t.rows:
            for cell in t.rows[0].cells:
                set_cell_shading(cell, "D9E2F3")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(10)
    d.save(path)


def build(code, out_name):
    page = COURSES / f"{code}-syllabus.html"
    m = ARTICLE.search(page.read_text(encoding="utf-8"))
    if not m:
        sys.exit(f"{page.name}: no <article class=\"syllabus-doc\"> found")
    out = HERE / out_name
    with tempfile.NamedTemporaryFile("w", suffix=".html", delete=False, encoding="utf-8") as tmp:
        tmp.write(to_word_html(m.group(1)))
    subprocess.run(
        ["pandoc", tmp.name, "-f", "html", "-t", "docx",
         "--reference-doc", str(HERE / "reference.docx"), "-o", str(out)],
        check=True,
    )
    Path(tmp.name).unlink()
    polish(out)
    print(f"{page.name} -> {out.relative_to(COURSES.parent)}")


if __name__ == "__main__":
    wanted = sys.argv[1:] or list(COURSES_TO_BUILD)
    for code in wanted:
        if code not in COURSES_TO_BUILD:
            sys.exit(f"unknown course {code}; choose from {', '.join(COURSES_TO_BUILD)}")
        build(code, COURSES_TO_BUILD[code])
